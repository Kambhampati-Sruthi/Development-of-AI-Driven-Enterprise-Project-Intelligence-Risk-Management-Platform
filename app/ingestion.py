
from io import BytesIO
from pathlib import Path
import json
import re
import xml.etree.ElementTree as ET
from html.parser import HTMLParser

import pandas as pd
from docx import Document
from pypdf import PdfReader
from pptx import Presentation


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".log", ".yaml", ".yml",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c",
    ".cpp", ".h", ".hpp", ".cs", ".go", ".rs", ".php",
    ".rb", ".swift", ".kt", ".sql", ".sh", ".bat", ".ps1",
    ".css", ".scss", ".ini", ".cfg", ".conf", ".toml",
    ".env.example", ".r", ".ipynb.txt",
}


class _HTMLTextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        text = data.strip()
        if text:
            self.parts.append(text)


def _record(text, source, **metadata):
    text = str(text).strip()
    if not text:
        return None

    return {
        "text": text,
        "source": source,
        **metadata,
    }


def extract_document(filename, content):
    """
    Extract readable text records from a supported file.

    filename: original filename, including its extension
    content: file bytes
    returns: list of dictionaries with text and source metadata
    """
    filename = Path(filename).name
    suffix = Path(filename).suffix.lower()
    source = filename
    records = []

    if not content:
        raise ValueError("The uploaded file is empty.")

    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))

            for page_number, page in enumerate(reader.pages, start=1):
                item = _record(
                    page.extract_text() or "",
                    source,
                    page=page_number,
                )
                if item:
                    records.append(item)

        elif suffix == ".docx":
            doc = Document(BytesIO(content))

            for index, paragraph in enumerate(doc.paragraphs, start=1):
                item = _record(
                    paragraph.text,
                    source,
                    paragraph=index,
                )
                if item:
                    records.append(item)

            for table_index, table in enumerate(doc.tables, start=1):
                for row_index, row in enumerate(table.rows, start=1):
                    text = " | ".join(
                        cell.text.strip() for cell in row.cells
                    )
                    item = _record(
                        text,
                        source,
                        table=table_index,
                        row=row_index,
                    )
                    if item:
                        records.append(item)

        elif suffix == ".csv":
            frame = pd.read_csv(BytesIO(content)).fillna("")

            for row_number, row in enumerate(
                frame.to_dict(orient="records"), start=2
            ):
                text = " | ".join(
                    f"{key}: {value}" for key, value in row.items()
                )
                item = _record(text, source, row=row_number)
                if item:
                    records.append(item)

        elif suffix == ".tsv":
            frame = pd.read_csv(
                BytesIO(content), sep="\t"
            ).fillna("")

            for row_number, row in enumerate(
                frame.to_dict(orient="records"), start=2
            ):
                text = " | ".join(
                    f"{key}: {value}" for key, value in row.items()
                )
                item = _record(text, source, row=row_number)
                if item:
                    records.append(item)

        elif suffix in {".xlsx", ".xls"}:
            sheets = pd.read_excel(
                BytesIO(content),
                sheet_name=None,
            )

            for sheet_name, frame in sheets.items():
                frame = frame.fillna("")

                for row_number, row in enumerate(
                    frame.to_dict(orient="records"), start=2
                ):
                    text = " | ".join(
                        f"{key}: {value}" for key, value in row.items()
                    )
                    item = _record(
                        f"Sheet: {sheet_name}\n{text}",
                        source,
                        sheet=sheet_name,
                        row=row_number,
                    )
                    if item:
                        records.append(item)

        elif suffix == ".json":
            parsed = json.loads(content.decode("utf-8-sig"))
            formatted = json.dumps(
                parsed,
                ensure_ascii=False,
                indent=2,
            )
            item = _record(formatted, source)
            if item:
                records.append(item)

        elif suffix in {".html", ".htm"}:
            parser = _HTMLTextParser()
            parser.feed(content.decode("utf-8-sig", errors="replace"))
            item = _record("\n".join(parser.parts), source)
            if item:
                records.append(item)

        elif suffix in {".xml", ".xsd", ".xsl"}:
            root = ET.fromstring(content)
            parts = []

            for element in root.iter():
                if element.text and element.text.strip():
                    tag = element.tag.split("}")[-1]
                    parts.append(f"{tag}: {element.text.strip()}")

                for key, value in element.attrib.items():
                    parts.append(f"{key}: {value}")

            item = _record("\n".join(parts), source)
            if item:
                records.append(item)

        elif suffix == ".pptx":
            presentation = Presentation(BytesIO(content))

            for slide_number, slide in enumerate(
                presentation.slides, start=1
            ):
                parts = []

                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        if shape.text.strip():
                            parts.append(shape.text.strip())

                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            parts.append(
                                " | ".join(
                                    cell.text.strip()
                                    for cell in row.cells
                                )
                            )

                item = _record(
                    "\n".join(parts),
                    source,
                    slide=slide_number,
                )
                if item:
                    records.append(item)

        elif suffix in TEXT_EXTENSIONS or suffix in {".txt", ".md"}:
            text = content.decode("utf-8-sig", errors="replace")
            item = _record(text, source)
            if item:
                records.append(item)

        else:
            raise ValueError(
                f"Unsupported file type: {suffix or '(no extension)'}. "
                "Supported formats include PDF, DOCX, CSV, TSV, TXT, "
                "MD, JSON, XLSX, XLS, PPTX, HTML, XML, YAML, logs, "
                "and common source-code files."
            )

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Could not read {filename}: {exc}"
        ) from exc

    if not records:
        raise ValueError(
            "No readable text was extracted. The file may be empty, "
            "contain only images, or require OCR."
        )

    return records
